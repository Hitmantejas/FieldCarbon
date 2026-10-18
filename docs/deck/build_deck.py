"""Fill the official AI Builder Cup submission template with the FieldCarbon deck.

  python docs/deck/build_deck.py "<path to Submission Template _ AI Builder Cup.pptx>"
      -> docs/deck/FieldCarbon_AI_Builder_Cup.pptx

Needs python-pptx, which is not a project dependency: run it from a scratch environment.
Text mirrors docs/deck/deck.md; every figure comes from the app's endpoints or evals/results/summary.md.
Diagrams are native shapes so they stay editable. Slides 2 (template note), 7 (optional wireframes)
and 15 (blank spare) are removed.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("FieldCarbon_AI_Builder_Cup.pptx")
FONT = "Google Sans Flex"
INK = RGBColor(0x20, 0x21, 0x24)
SOFT = RGBColor(0x5F, 0x63, 0x68)
MOSS = RGBColor(0x2F, 0x6B, 0x3F)
BLUE = RGBColor(0x1A, 0x73, 0xE8)
GREEN = RGBColor(0x18, 0x80, 0x38)
TODO = RGBColor(0xC5, 0x22, 0x1F)
FILL = RGBColor(0xF1, 0xF3, 0xF4)
LINE = RGBColor(0xDA, 0xDC, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TOP = 1.35  # inches: below the template's slide title

REPO = "https://github.com/Hitmantejas/FieldCarbon"
LIVE = "https://fieldcarbon-frontend-818326148498.us-central1.run.app"


# ---------- text helpers ----------

def _runs(p, text: str, size: float, color: RGBColor, bold: bool = False):
    """Add text to a paragraph; **double stars** mark bold segments."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or i % 2 == 1


def _bullet(p, level: int = 0):
    pPr = p._p.get_or_add_pPr()
    indent = Inches(0.16)
    pPr.set("marL", str(int(indent * (level + 1))))
    pPr.set("indent", str(-int(indent)))
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def write(tf, items, size=12, color=INK, space=4):
    """items: str (paragraph), ("-", str) bullet, ("h", str) heading, ("s", str) small grey note."""
    tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        kind, text = ("p", it) if isinstance(it, str) else it
        p.space_after = Pt(space)
        if kind == "h":
            _runs(p, text, size + 1, MOSS, bold=True)
        elif kind == "s":
            _runs(p, text, size - 2, SOFT)
        elif kind == "todo":
            _runs(p, text, size, TODO, bold=True)
        else:
            _runs(p, text, size, color)
            if kind == "-":
                _bullet(p)
    return tf


def textbox(slide, x, y, w, h, items, size=12, color=INK, anchor=MSO_ANCHOR.TOP, align=None, space=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    write(tf, items, size, color, space)
    if align:
        for p in tf.paragraphs:
            p.alignment = align
    return tb


def box(slide, x, y, w, h, items, size=10, fill=FILL, line=LINE, color=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        anchor=MSO_ANCHOR.TOP, align=None, dash=False, space=2):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(1.25)
    if dash:
        s.line.dash_style = 4  # dash
    s.shadow.inherit = False
    tf = s.text_frame
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.08))
    for side in ("margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.05))
    write(tf, items, size, color, space)
    for p in tf.paragraphs:
        p.alignment = align or PP_ALIGN.LEFT
    return s


def arrow(slide, x1, y1, x2, y2, both=False, color=SOFT):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle"}))
    return c


def tag(slide, x, y, text, color):
    s = box(slide, x, y, 0.95, 0.22, [text], size=8, fill=color, line=color, color=WHITE,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    s.adjustments[0] = 0.5
    s.text_frame.margin_top = s.text_frame.margin_bottom = 0
    s.text_frame.paragraphs[0].runs[0].font.bold = True
    return s


def table(slide, x, y, w, rows, widths, size=10, row_h=0.28, header=True):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shp.table
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            head = header and i == 0
            cell.fill.fore_color.rgb = INK if head else (FILL if i % 2 == 0 else WHITE)
            tf = cell.text_frame
            tf.text = ""
            _runs(tf.paragraphs[0], val, size, WHITE if head else INK, bold=head)
    return t


def title_of(slide):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())


def set_after(paragraph, text, size=None, bold=False, color=INK, link=None):
    """Append a run to an existing template paragraph, keeping its font."""
    base = paragraph.runs[0].font if paragraph.runs else None
    r = paragraph.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size) if size else (base.size if base else None)
    r.font.bold = bold
    r.font.color.rgb = color
    if link:
        r.hyperlink.address = link
    return r


def drop(prs, indexes):
    ids = prs.slides._sldIdLst
    for i in sorted(indexes, reverse=True):
        sld = ids[i]
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)


# ---------- slides ----------

def team(s):
    paras = title_of(s).text_frame.paragraphs
    texts = [p.text for p in paras]
    set_after(paras[texts.index("Team name: ")], "Valtechies", bold=True)
    set_after(paras[texts.index("Team leader name: ")], "Tejas Bhorade", bold=True)
    set_after(paras[texts.index("Problem Statement: ")],
              "Sustainability & Social Impact — fuel and CO2e for farm-machinery fleets, with honest confidence",
              size=14, bold=True)


def brief(s):
    textbox(s, 0.3, TOP, 9.4, 0.62, [
        "**FieldCarbon: where a farm-machinery fleet wastes fuel, what to change, and how sure we are.** "
        "Every figure traceable."], size=15)
    box(s, 0.3, 2.05, 4.6, 2.45, [
        ("h", "The problem"),
        ("-", "Agricultural contractors and co-operatives run tractors across many client farms."),
        ("-", "Buyers now ask them for fuel and CO2e **per farm and per job**."),
        ("-", "Fuel is bought in bulk, dispensed from one depot tank and burned across dozens of clients; "
              "the paperwork is on paper or in people's heads."),
        ("-", "Carbon tools give one number with **no error bar** and no next step."),
    ], size=11, space=4)
    box(s, 5.1, 2.05, 4.6, 2.45, [
        ("h", "FieldCarbon"),
        ("-", "Turns telemetry, fuel logs, **photos of delivery notes and job sheets**, and **voice notes** into "
              "fuel, cost and CO2e by site, machine, job, field and client."),
        ("-", "Finds waste and tests its causes; ranks what to change by **strength of evidence**."),
        ("-", "Proves savings before and after; writes a **traceable statement per client farm**."),
    ], size=11, space=4)
    box(s, 0.3, 4.62, 9.4, 0.42, [
        "**Gemini reads and explains.  Tested code does every calculation.  An auditor checks every number Gemini writes.**"],
        size=12, fill=INK, line=INK, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def opportunities(s):
    tb = title_of(s)
    tf = tb.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    tb.height = Inches(0.6)
    cols = [
        ("How is it different?", [
            ("-", "Other tools give **a number**. FieldCarbon says **what to change, how sure we are, and what data "
                  "would make us surer**."),
            ("-", "Every figure carries a label as text: **MEASURED, ALLOCATED or ESTIMATED**, and opens into its "
                  "sources, emission factor and calculation."),
            ("-", "Savings are **ranges** from published research, never a single promise."),
        ]),
        ("How does it solve the problem?", [
            ("-", "Uses data contractors already have: telemetry CSVs, a fuel log, **photos of paperwork**, voice notes."),
            ("-", "Splits fleet fuel to **each client farm and job**, so a contractor can answer a buyer's request."),
            ("-", "Checks deliveries vs dispenses vs fuel burned, and **tests the causes** of each gap."),
            ("-", "Ranks only changes that fit this fleet."),
        ]),
        ("USP", [
            ("-", "**Uncertainty is the product:** labels on 100% of figures, and \"what would improve this answer?\" "
                  "names the one missing input that most narrows a range."),
            ("-", "**Gemini never does the maths:** an auditor traces every number in an AI answer to a tool output."),
            ("-", "**Proof, not promises:** savings judged against machines that did not change."),
        ]),
    ]
    for i, (head, items) in enumerate(cols):
        box(s, 0.3 + i * 3.17, TOP, 3.06, 3.15, [("h", head), *items], size=10.5, space=5)


FEATURES = [
    ("Carbon ledger", "Tonnes emitted, a bounded range that could be avoided, tonnes proven avoided, goal check."),
    ("Week summary", "The three things to look at first, chosen by fixed rules in code."),
    ("Sites, machines, jobs", "The same job compared across machines: L/ha, L/h, idle share."),
    ("Client farms", "Fuel and CO2e for every client farm the fleet works for."),
    ("Fuel control", "Depot deliveries vs dispenses vs fuel burned, month by month."),
    ("Delivery-note reading", "Gemini reads photos (also live upload); code checks them against the fuel log."),
    ("Anomaly diagnosis", "Code tests a fixed menu of causes; Gemini suggests; a person decides."),
    ("Voice notes & job sheets", "Gemini hears machine faults and reads handwritten workshop sheets."),
    ("Ranked ideas & what-if", "8 interventions with litres, CO2e, cost and payback ranges, and their evidence."),
    ("Before-and-after proof", "Logged changes judged against machines that did not change."),
    ("Fleet analyst agent", "Ask in plain language; shows its tools; every number traced."),
    ("Client reports", "One-page statement in English, German, Spanish or Hindi; print or PDF."),
]


def features(s):
    w, h, gx, gy = 2.27, 1.2, 0.11, 0.1
    for i, (name, text) in enumerate(FEATURES):
        r, c = divmod(i, 4)
        box(s, 0.3 + c * (w + gx), TOP + r * (h + gy), w, h, [("h", name), text], size=10, space=3)


STEPS = [
    ("1  Capture", "Telemetry CSV, fuel log, photos of delivery notes and job sheets, voice notes.", "You"),
    ("2  Read", "Gemini 3.5 Flash turns photos and audio into a fixed JSON schema; text on a document is never followed.", "Gemini"),
    ("3  Check", "Rules R1–R8; compared with the fuel log: matches, differs, not in log, keyed twice, needs review.", "Code"),
    ("4  Calculate", "Litres × UK GHG 2025 factor (well-to-wheel) in BigQuery; every figure labelled.", "Code"),
    ("5  Find & diagnose", "Fuel gaps, frozen recorders, unusual jobs; causes tested against the data.", "Code"),
    ("6  Act", "Ideas ranked by evidence; savings as ranges; \"what would improve this answer?\"", "Code"),
    ("7  Prove", "Before-and-after against machines that did not change.", "Code"),
    ("8  Report & ask", "Client statement in 4 languages; the analyst agent answers questions.", "Gemini + auditor"),
]
TAGS = {"You": SOFT, "Gemini": BLUE, "Code": GREEN, "Gemini + auditor": BLUE}


def process(s):
    w, h, gx = 2.05, 1.35, 0.4
    for i, (head, text, who) in enumerate(STEPS):
        r, c = divmod(i, 4)
        x, y = 0.3 + c * (w + gx), TOP + 0.05 + r * (h + 0.35)
        box(s, x, y, w, h, [("h", head), text], size=9.5, space=2)
        t = tag(s, x + w - 1.0, y + h - 0.28, who, TAGS[who])
        if who == "Gemini + auditor":
            t.width = Inches(1.15)
            t.left = Inches(x + w - 1.2)
        if c < 3:
            arrow(s, x + w + 0.03, y + h / 2, x + w + gx - 0.03, y + h / 2)
    # row 1 end -> row 2 start: down, back along the gap, down
    mid, x4, x1 = TOP + 0.05 + h + 0.175, 0.3 + 3 * (w + gx) + w / 2, 0.3 + w / 2
    for a, b in (((x4, TOP + 0.05 + h), (x4, mid)), ((x4, mid), (x1, mid))):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1]))
        ln.line.color.rgb, ln.line.width = SOFT, Pt(1.5)
    arrow(s, x1, mid, x1, TOP + 0.05 + h + 0.35)
    textbox(s, 0.3, 4.62, 9.4, 0.45, [
        "**Use case:** the fleet manager asks \"Where did we waste fuel this season, and what should we change first?\" → "
        "the agent calls the tools → an answer whose every number is traced → the manager prints each client's statement."],
        size=10)


def architecture(s):
    y0 = TOP
    inputs = [
        ("Tractor telemetry", "3 public datasets (CAN bus, GNSS, fuel rate)"),
        ("Paperwork photos", "Delivery notes, receipts, workshop sheets"),
        ("Reference tables", "Emission factors, fuel-use rates, 8 interventions, prices"),
    ]
    for i, (h, t) in enumerate(inputs):
        box(s, 0.3, y0 + 0.2 + i * 1.2, 1.75, 0.95, [("h", h), t], size=9)
    frame = box(s, 2.35, y0, 5.6, 3.5, [("h", "Google Cloud  ·  us-central1")], size=10, fill=None, line=BLUE, dash=True)
    frame.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
    box(s, 2.55, y0 + 0.45, 1.6, 0.8, [("h", "Cloud Storage"), "raw archives, images"], size=9, fill=WHITE)
    box(s, 2.55, y0 + 1.75, 1.6, 1.5, [("h", "BigQuery"), "dataset fieldcarbon", "tables, SQL views (litres × factor), GIS for fields"],
        size=9, fill=WHITE)
    box(s, 4.5, y0 + 0.45, 1.75, 0.85, [("h", "Vertex AI"), "Gemini 3.5 Flash", "(global endpoint)"], size=9, fill=WHITE, line=BLUE)
    box(s, 4.5, y0 + 1.75, 1.75, 1.5, [("h", "Cloud Run: backend"), ("-", "FastAPI"), ("-", "deterministic tools"),
                                        ("-", "ADK analyst agent"), ("-", "auditor")], size=9, fill=WHITE)
    box(s, 6.55, y0 + 1.75, 1.25, 1.0, [("h", "Cloud Run: frontend"), "React app"], size=9, fill=WHITE)
    box(s, 8.2, y0 + 1.75, 1.5, 1.0, [("h", "Fleet manager"), "browser; uploads photos and voice notes"], size=9)
    # telemetry, photos -> storage; reference -> BigQuery
    arrow(s, 2.05, y0 + 0.67, 2.55, y0 + 0.75)
    arrow(s, 2.05, y0 + 1.87, 2.55, y0 + 1.05)
    arrow(s, 2.05, y0 + 3.07, 2.55, y0 + 2.75)
    arrow(s, 3.35, y0 + 1.25, 3.35, y0 + 1.75)  # storage -> BigQuery
    arrow(s, 4.15, y0 + 2.5, 4.5, y0 + 2.5, both=True)  # BigQuery <-> backend
    arrow(s, 5.37, y0 + 1.3, 5.37, y0 + 1.75, both=True, color=BLUE)  # Gemini <-> backend
    arrow(s, 6.25, y0 + 2.25, 6.55, y0 + 2.25, both=True)  # backend <-> frontend
    arrow(s, 7.8, y0 + 2.25, 8.2, y0 + 2.25, both=True)  # frontend <-> user
    textbox(s, 0.3, y0 + 3.6, 9.4, 0.4, [
        "Gemini reads documents and audio, plans tool calls and explains. **Every number comes from the tools**; the auditor "
        "traces each number in an answer. Only synthetic documents and demo questions are sent to the model."], size=9.5, color=SOFT)


def technologies(s):
    rows = [
        ["Layer", "Technology"],
        ["AI model", "**Gemini 3.5 Flash on Vertex AI** (multimodal: images and audio), global endpoint"],
        ["Agent", "**Google Agent Development Kit (ADK)**, tool calling; deterministic auditor for faithfulness"],
        ["Data", "**BigQuery** (tables, SQL views, BigQuery GIS), **Cloud Storage**"],
        ["Hosting", "**Cloud Run** (frontend and backend), Cloud Build, Artifact Registry"],
        ["Backend", "Python, FastAPI, pytest (every calculation unit-tested); eval harness"],
        ["Frontend", "React, Vite, Tailwind CSS; WCAG contrast, labels as text, 4 languages"],
        ["Reference data", "UK GHG Conversion Factors 2025 · Iowa State A3-27 and Purdue AE-110 fuel-use rates · "
                           "Defra red diesel prices · peer-reviewed and extension studies"],
    ]
    table(s, 0.3, TOP + 0.05, 9.4, rows, [1.7, 7.7], size=11, row_h=0.46)


def cost(s):
    cards = [
        ("Cloud Run", ["Scales to zero: pay only while the app is in use.", "Two small services (frontend, backend)."]),
        ("BigQuery and Storage", ["The whole fleet is a few MB.", "Queries stay in the free tier at this size."]),
        ("Gemini (measured tokens)", ["About **3,000 tokens** per document read", "About **1,600** per voice note",
                                      "About **3,500** per client statement", "About **17,400** per analyst question"]),
    ]
    for i, (h, items) in enumerate(cards):
        box(s, 0.3 + i * 3.17, TOP, 3.06, 1.45, [("h", h), *[("-", t) for t in items]], size=10.5, space=3)
    box(s, 0.3, TOP + 1.6, 9.4, 1.15, [
        ("h", "Example month for a contractor"),
        "50 delivery notes + 20 voice notes + 10 client statements + 100 analyst questions "
        "≈ **1.95 million tokens** (150k + 32k + 35k + 1.74M).",
        ("todo", "TODO: × current Vertex AI price for gemini-3.5-flash = $ per month."),
        ("s", "Token counts are the ones Vertex AI reported in our live runs; analyst questions dominate the cost."),
    ], size=11, space=4)


def snapshots(s):
    shots = ["Overview: carbon ledger and week summary", "Fuel control: a delivery note read and checked",
             "Diagnosis: causes supported / not supported", "Opportunities: ranked ideas with ranges",
             "Ask the analyst: \"Numbers checked by code\"", "Reports: client statement with labels"]
    w, h = 3.06, 1.78
    for i, name in enumerate(shots):
        r, c = divmod(i, 3)
        box(s, 0.3 + c * (w + 0.11), TOP + r * (h + 0.12), w, h, [("todo", "Insert screenshot"), name], size=10,
            fill=WHITE, dash=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)


def performance(s):
    kpis = [("16,492 L", "fuel (allocated)"), ("55.8 t", "CO2e (allocated)"), ("713", "jobs"),
            ("1,800 ha", "worked"), ("1.1–3.3 t/yr", "cut (estimated)")]
    for i, (v, t) in enumerate(kpis):
        b = box(s, 0.3 + i * 1.9, TOP - 0.05, 1.8, 0.62, [v, ("s", t)], size=15, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
                space=0)
        b.text_frame.paragraphs[0].runs[0].font.bold = True
    rows = [
        ["Evaluation (python -m evals.run_all)", "Result"],
        ["Tool eval cases · figures with a confidence label and sources", "72 of 72 · 100% of 918"],
        ["Analyst agent, 43 live questions: right tools / every number traced / ideas cite source", "100% / 100% / 100%"],
        ["Agent refusals (ranking operators, farming advice, forecasts, guarantees, injected rules)", "7 of 7"],
        ["Gemini reading, every field: 24 receipts, 14 delivery notes, 36 workshop sheets", "100%"],
        ["Voice notes, 8 recorded clips: machine and issue / expected status", "8 of 8 / 8 of 8"],
        ["Planted fuel problems found / top supported cause is the planted one", "3 of 3 / 3 of 3"],
        ["Before-and-after, 150 random trials: false saving claimed / 90% range holds truth", "9% / 91–97%"],
        ["Receipts-only estimate vs telemetry truth (job types inside the range)", "8 of 12"],
    ]
    table(s, 0.3, TOP + 0.7, 9.4, rows, [7.3, 2.1], size=9.5, row_h=0.29)
    textbox(s, 0.3, TOP + 3.38, 9.4, 0.4, [
        ("s", "Honest notes: 17 of 43 first agent answers needed one automatic rewrite. Document images are synthetic renders, "
              "so 100% shows the pipeline works, not every real receipt. Demo data: illustrative.")], size=11)


def future(s):
    box(s, 0.3, TOP, 4.6, 2.2, [
        ("h", "Next"),
        ("-", "Live fleet telemetry exports (CSV today, no brand lock-in)."),
        ("-", "Co-operative and advisor view across many contractors."),
        ("-", "Field-level soil and terrain at scale; alerts when a machine drifts from its peers."),
        ("-", "Production: scheduled BigQuery loads, data-quality checks, layered ingestion."),
    ], size=11, space=5)
    box(s, 5.1, TOP, 4.6, 2.2, [
        ("h", "Responsible by design"),
        ("-", "Operators are never named, ranked or compared."),
        ("-", "Uploads are checked by their bytes and never stored; no coordinates published."),
        ("-", "Prompts ignore instructions written on documents or spoken in recordings."),
        ("-", "Reports say \"traceable\", never \"verified\"; savings are never guaranteed."),
    ], size=11, space=5)
    textbox(s, 0.3, TOP + 2.35, 9.4, 0.8, [
        ("s", "Data credits: Götz et al. (2025), TUM Agricultural Load Cycles, CC BY 4.0 · Rossi & Alemanno, tractor "
              "performances in vineyard operations, CC BY 4.0 · Dana, Ozener & Tellioglu, Tümosan tillage CAN bus data, "
              "CC BY 4.0 · UK DESNZ GHG Conversion Factors 2025, Open Government Licence v3.0. Brands removed; data aggregated.")],
        size=11)


def links(s):
    paras = title_of(s).text_frame.paragraphs
    for p in paras:
        if p.text == "GitHub Public Repository":
            set_after(p, ": ", size=13)
            set_after(p, REPO, size=13, color=BLUE, link=REPO)
        elif p.text == "Demo Video Link (3 Minutes)":
            set_after(p, ": ", size=13)
            set_after(p, "TODO", size=13, bold=True, color=TODO)
        elif p.text == "Final Product Link":
            set_after(p, ": ", size=13)
            set_after(p, LIVE, size=13, color=BLUE, link=LIVE)


def build(template: str) -> Path:
    prs = Presentation(template)
    sl = list(prs.slides)
    team(sl[0])
    brief(sl[2])
    opportunities(sl[3])
    features(sl[4])
    process(sl[5])
    architecture(sl[7])
    technologies(sl[8])
    cost(sl[9])
    snapshots(sl[10])
    performance(sl[11])
    future(sl[12])
    links(sl[13])
    drop(prs, [1, 6, 14])
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build(sys.argv[1]))
